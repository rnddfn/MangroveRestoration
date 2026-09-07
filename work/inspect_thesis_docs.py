import argparse
from pathlib import Path
from docx import Document


ROOT = Path(r"C:\Users\lenov\Mangrove\skripsi\draft")

parser = argparse.ArgumentParser()
parser.add_argument("--name")
parser.add_argument("--start", type=int, default=0)
parser.add_argument("--end", type=int, default=10_000)
args = parser.parse_args()


for path in sorted(ROOT.glob("*.docx")):
    if args.name and path.name != args.name:
        continue
    doc = Document(path)
    print(f"\n=== {path.name} ===")
    print(
        f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} "
        f"sections={len(doc.sections)} inline_shapes={len(doc.inline_shapes)}"
    )
    for index, paragraph in enumerate(doc.paragraphs):
        if not (args.start <= index < args.end):
            continue
        text = paragraph.text.strip()
        if text:
            print(index, repr(paragraph.style.name), repr(text[:1000]))
    for index, table in enumerate(doc.tables):
        print(f"TABLE {index}: rows={len(table.rows)} cols={len(table.columns)}")
        for row in table.rows[:12]:
            print(" | ".join(cell.text.replace("\n", " / ")[:250] for cell in row.cells))
